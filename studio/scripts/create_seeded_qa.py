#!/usr/bin/env python3
import argparse,json
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches,Pt
from pptx.enum.text import MSO_AUTO_SIZE

p=argparse.ArgumentParser();p.add_argument('--out',required=True);a=p.parse_args();out=Path(a.out);out.mkdir(parents=True,exist_ok=True)
r=Presentation();r.slide_width=Inches(13.333333);r.slide_height=Inches(7.5);s=r.slides.add_slide(r.slide_layouts[6])
def box(text,x,y,w,h,size):
 b=s.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h));b.text_frame.word_wrap=True;b.text_frame.auto_size=MSO_AUTO_SIZE.NONE
 para=b.text_frame.paragraphs[0];para.text=text;para.font.name='DejaVu Sans';para.font.size=Pt(size)
box('QA SEED / INTENTIONAL DEFECTS / DO NOT USE',.6,.5,12,1,30)
box('Все компании удваивают прибыль. В группе одновременно 12 и 15 наблюдений.',.7,2,11.8,1,22)
box('Источник: https://example.org — общая страница без подтверждающего фрагмента.',.7,3.5,11.8,.7,14)
box(' '.join(f'Overflow sentinel {i:03}' for i in range(220)),8,5.4,4,.35,22)
r.save(out/'error-seeded.pptx')
ledger={'sources':[{'id':'weak','url':'https://example.org','title':'Synthetic weak-reference fixture','publisher':'Test fixture','accessed':'2026-10-22','locator':'Generic landing page, no relevant evidence','support':'No supporting passage; intentionally weak fixture','access_state':'not_opened','evidence_role':'discovery_only'}],
 'claims':[{'id':'unsupported','kind':'fact','text':'All companies double profit.'},{'id':'weakly-cited','kind':'fact','text':'A generic homepage proves the profit claim.','source_ids':['weak']}], 'calculations':[]}
for i,n in enumerate([12,15]):ledger['claims'].append({'id':f'wrong-number-{i}','kind':'synthetic','text':f'The same group contains {n} observations.','source_ids':['weak'],'disclosure':'Intentional QA fixture','metric':'n','value':n,'unit':'observations','period':'test','geography':'fictional','scope':'same group'})
(out/'error-seeded.json').write_text(json.dumps(ledger,ensure_ascii=False,indent=2));print(out)
