#!/usr/bin/env python3
"""Synthetic editable workbook and a three-page research note with traceable data."""
import argparse,json,hashlib
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font,PatternFill,Alignment,Border,Side
from openpyxl.chart import BarChart,Reference
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt

def workbook(data,out):
    w=Workbook();s=w.active;s.title='Raw'
    s.append(['ID','Контроль, мин','Поддержка, мин'])
    for i,(left,right) in enumerate(zip(data['control'],data['assisted']),1):s.append([i,left,right])
    s.freeze_panes='B2';s.auto_filter.ref='A1:C7'
    for col in ['A','B','C']:s.column_dimensions[col].width=24
    summary=w.create_sheet('Analysis')
    summary['A1']='ИСКУССТВЕННЫЕ ВЫБОРКИ / 22.10.2026';summary.merge_cells('A1:H1')
    for col,val in zip(['A3','B3','C3'],['Метрика','Контроль','Поддержка']):summary[col]=val
    for row,label,fn in [(4,'Число наблюдений','COUNT'),(5,'Медиана, мин','MEDIAN'),(6,'Среднее, мин','AVERAGE'),(7,'Минимум, мин','MIN'),(8,'Максимум, мин','MAX'),(9,'Сумма, мин','SUM')]:
        summary.cell(row,1,label)
        for col,rawcol in [(2,'B'),(3,'C')]:summary.cell(row,col,f'={fn}(Raw!{rawcol}2:{rawcol}7)')
    summary['A11']='Разница медиан, мин';summary['B11']='=B5-C5'
    summary['A12']='Относительная разница';summary['B12']='=IFERROR(B11/B5,0)';summary['B12'].number_format='0.0%'
    for row,vals in enumerate([['Группа','Медиана, мин'],['Контроль','=B5'],['Поддержка','=C5']],16):
        for col,v in enumerate(vals,1):summary.cell(row,col,v)
    summary['A21']='Источник: studio/examples/research-data.json';summary.merge_cells('A21:H21')
    summary['A22']='Все значения вымышленные; относительная разница не является причинным эффектом.';summary.merge_cells('A22:H22')
    summary.column_dimensions['A'].width=30
    for col in ['B','C']:summary.column_dimensions[col].width=17
    summary.column_dimensions['D'].width=4
    for col in ['E','F','G','H','I','J','K','L']:summary.column_dimensions[col].width=10
    summary.freeze_panes='B4';summary.sheet_view.showGridLines=False
    chart=BarChart();chart.type='col';chart.style=10;chart.title='Медианное время';chart.y_axis.title='Минуты';chart.y_axis.scaling.min=0;chart.y_axis.scaling.max=50;chart.height=8.7;chart.width=14
    chart.add_data(Reference(summary,min_col=2,min_row=16,max_row=18),titles_from_data=True)
    chart.set_categories(Reference(summary,min_col=1,min_row=17,max_row=18));chart.legend=None
    chart.series[0].graphicalProperties.solidFill='356D93';summary.add_chart(chart,'E3')
    summary.print_area='A1:L23';summary.page_setup.orientation='landscape';summary.page_setup.paperSize=summary.PAPERSIZE_A4;summary.page_setup.fitToWidth=1;summary.page_setup.fitToHeight=1;summary.sheet_properties.pageSetUpPr.fitToPage=True
    sources=w.create_sheet('Sources');sources.append(['ID','Источник','Область','Статус'])
    sources.append([data['id'],'file:'+data['source'],'Все 12 наблюдений','Синтетический тест, не эмпирические данные'])
    for col in ['A','B','C','D']:sources.column_dimensions[col].width=35
    sources.row_dimensions[2].height=44;sources.print_area='A1:D3';sources.page_setup.orientation='landscape';sources.page_setup.paperSize=sources.PAPERSIZE_A4;sources.page_setup.fitToWidth=1;sources.page_setup.fitToHeight=1;sources.sheet_properties.pageSetUpPr.fitToPage=True
    s.print_area='A1:C9';s.page_setup.fitToWidth=1;s.page_setup.fitToHeight=1;s.sheet_properties.pageSetUpPr.fitToPage=True
    for sheet in w:
        for row in sheet:
            for cell in row:
                cell.font=Font(name='Inter',size=11,color='253A4A');cell.alignment=Alignment(vertical='center',wrap_text=True)
        for row in ([1] if sheet!=summary else [3,16]):
            for cell in sheet[row]:cell.fill=PatternFill('solid',fgColor='EAF0F2');cell.font=Font(name='Inter',size=11,bold=True,color='253A4A')
        sheet.sheet_view.showGridLines=False
    for row in range(4,10):
        for col in [2,3]:summary.cell(row,col).number_format='0.0' if row in [5,6] else '0'
    summary['A1'].font=Font(name='Inter',size=15,bold=True,color='253A4A')
    w.save(out/'research-workbook.xlsx')

def plot(data,out):
    fig,ax=plt.subplots(figsize=(7.2,2.8));fig.patch.set_facecolor('#FAF8F3');ax.set_facecolor('#FAF8F3')
    for vals,y,color,label in [(data['control'],1,'#356D93','Контроль'),(data['assisted'],0,'#C1784C','Поддержка')]:
        ax.scatter(vals,[y+(.08 if i%2 else -.08) for i in range(len(vals))],s=40,color=color,zorder=3)
        sorted_vals=sorted(vals);median=(sorted_vals[2]+sorted_vals[3])/2
        ax.plot([min(vals),max(vals)],[y,y],color=color,lw=2)
        ax.plot([median,median],[y-.24,y+.24],color=color,lw=2)
    ax.set_xlim(0,50);ax.set_ylim(-.5,1.5);ax.set_yticks([0,1],['Поддержка','Контроль']);ax.set_xlabel('Время, минуты · искусственные значения');ax.grid(axis='x',color='#CED9DF',lw=.6);ax.set_axisbelow(True)
    for spine in ax.spines.values():spine.set_visible(False)
    fig.tight_layout();fig.savefig(out/'observations.png',dpi=180);plt.close(fig)

def shade(cell,color):
    pr=cell._tc.get_or_add_tcPr();sh=OxmlElement('w:shd');sh.set(qn('w:fill'),color);pr.append(sh)

def document(data,analysis,out):
    d=Document();section=d.sections[0];section.top_margin=Inches(.65);section.bottom_margin=Inches(.6);section.left_margin=Inches(.7);section.right_margin=Inches(.7)
    normal=d.styles['Normal'];normal.font.name='Inter';normal.font.size=Pt(10.5);normal.paragraph_format.space_after=Pt(7)
    for style,size in [('Title',28),('Heading 1',20),('Heading 2',13)]:
        st=d.styles[style];st.font.name='Lora';st.font.size=Pt(size);st.font.color.rgb=RGBColor.from_string('253A4A')
    # Remove inherited theme pointers that LibreOffice otherwise resolves to Carlito.
    for style_name in ['Normal','Title','Heading 1','Heading 2','Subtitle','Caption','Header','Footer','List Bullet','Title Char','Heading 1 Char','Heading 2 Char']:
        st=d.styles[style_name];family='Lora' if style_name in ['Title','Heading 1','Heading 2','Title Char','Heading 1 Char','Heading 2 Char'] else 'Inter'
        st.font.name=family
        fonts=st._element.get_or_add_rPr().find(qn('w:rFonts'))
        for key in list(fonts.attrib):
            if 'theme' in key.lower():del fonts.attrib[key]
        for key in ['ascii','hAnsi','eastAsia','cs']:fonts.set(qn('w:'+key),family)
        color=st._element.get_or_add_rPr().find(qn('w:color'))
        if color is not None:
            for key in list(color.attrib):
                if 'theme' in key.lower():del color.attrib[key]
            color.set(qn('w:val'),'526878' if style_name in ['Caption','Subtitle','Header','Footer'] else '253A4A')
        pr=st._element.find(qn('w:pPr'))
        if pr is not None:
            border=pr.find(qn('w:pBdr'))
            if border is not None:pr.remove(border)
    header=section.header.paragraphs[0];header.text='INTTRADE / RESEARCH NOTE 01 / 22.10.2026';header.style='Caption'
    footer=section.footer.paragraphs[0];footer.text='Синтетический пример · '
    field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)
    d.add_heading('Разница есть.\nПричина не установлена.',0)
    d.add_paragraph('Время подготовки отчёта в двух искусственных выборках',style='Subtitle')
    d.add_paragraph('Две выборки по шесть значений созданы для проверки анализа, формул, диаграмм и качества документов. Это не результаты наблюдения за людьми или компаниями. Все значения выражены в минутах.')
    d.add_heading('Описательный результат',1)
    d.add_paragraph('Медианное время составляет 41,5 минуты в группе «Контроль» и 30,5 минуты в группе «Поддержка». Разница равна 11 минутам. Названия групп условные и не обозначают проведённый эксперимент.')
    table=d.add_table(rows=1,cols=4);table.style='Light Shading Accent 1'
    for c,v in zip(table.rows[0].cells,['Группа','n','Медиана','Диапазон']):c.text=v;shade(c,'EAF0F2')
    for vals in [('Контроль','6','41,5 мин','39–45 мин'),('Поддержка','6','30,5 мин','28–34 мин')]:
        for c,v in zip(table.add_row().cells,vals):c.text=v
    d.add_picture(str(out/'observations.png'),width=Inches(6.6))
    d.add_paragraph('Все 12 точек, диапазоны и медианы показаны на одинаковой шкале от 0 до 50 минут.',style='Caption')
    d.add_paragraph('Источник: studio/examples/research-data.json. Данные вымышленные; рисунок показывает только свойства этого набора.',style='Caption')
    d.add_page_break();d.add_heading('Методология и исходные данные',0)
    d.add_paragraph('Единица наблюдения — время подготовки одного условного отчёта. Выборки заданы независимо. Номера строк нужны для воспроизведения таблицы и не означают пары участников. Дата выпуска примера — 22 октября 2026 года; периода реального наблюдения у него нет.')
    table=d.add_table(rows=1,cols=3);table.style='Light Shading Accent 1'
    for c,v in zip(table.rows[0].cells,['Строка','Контроль, мин','Поддержка, мин']):c.text=v
    for i,(x,y) in enumerate(zip(data['control'],data['assisted']),1):
        for c,v in zip(table.add_row().cells,[str(i),str(x),str(y)]):c.text=v
    d.add_heading('Как получены медианы',1)
    d.add_paragraph('После сортировки контрольная выборка имеет вид 39, 40, 41, 42, 44, 45. Медиана равна (41 + 42) / 2 = 41,5. Во второй выборке 28, 29, 30, 31, 32, 34 медиана равна (30 + 31) / 2 = 30,5. Разница: 41,5 − 30,5 = 11 минут.')
    d.add_paragraph('Отношение разницы к контрольной медиане составляет около 26,5%. Эта арифметическая характеристика вымышленных данных не является оценкой реального повышения эффективности.')
    d.add_heading('Воспроизводимость',1)
    d.add_paragraph('В XLSX лист Raw содержит исходные значения, Analysis — редактируемые формулы COUNT, MEDIAN, AVERAGE, MIN, MAX и SUM, а Sources — происхождение набора. Формулы пересчитаны офисным движком; результаты отдельно сопоставлены с исходными рядами.')
    d.add_page_break();d.add_heading('Границы выводов',0)
    d.add_paragraph('В этом примере не было отбора из популяции, рандомизации или измерения реальной работы. Нельзя оценить причинный эффект, устойчивость результата или его применимость к другой задаче. Малый размер вымышленных выборок не исправляется более точным оформлением диаграммы.')
    d.add_heading('Что проверять в реальном исследовании',1)
    for v in ['Сопоставимость заданий, опыта участников и условий выполнения.','Правила измерения времени, пропуски и возможные систематические ошибки.','Дизайн контрольного сравнения и необходимый размер выборки.','Даты, первоисточники, определения метрик и альтернативные объяснения.']:
        d.add_paragraph(v,style='List Bullet')
    d.add_heading('Источники и комплект файлов',1)
    d.add_paragraph('Единственный источник чисел — локальный синтетический набор research-data.json. Внешние источники не использованы для подтверждения коммерческих или рыночных утверждений: таких утверждений в записке нет.')
    d.add_paragraph('Комплект: исходный JSON, расчёт analysis.json, редактируемый research-workbook.xlsx, исследовательская презентация PPTX, этот DOCX и его фактический PDF-экспорт. Визуальные и числовые проверки перечислены отдельно; автоматическая проверка не заменяет оценку смысловой корректности.')
    digest=hashlib.sha256(json.dumps(data,sort_keys=True).encode()).hexdigest()
    d.add_paragraph('Идентификатор набора: '+data['id']+'. Контрольный SHA-256 нормализованного JSON: '+digest,style='Caption')
    d.core_properties.title='Исследовательская записка — синтетический пример';d.core_properties.author='IntTrade Research & Design Studio'
    d.save(out/'research-note.docx')

if __name__=='__main__':
    p=argparse.ArgumentParser();p.add_argument('--analysis',required=True);p.add_argument('--out',required=True);a=p.parse_args()
    out=Path(a.out);out.mkdir(parents=True,exist_ok=True)
    data=json.loads((Path(__file__).resolve().parents[1]/'examples/research-data.json').read_text());analysis=json.loads(Path(a.analysis).read_text())
    workbook(data,out);plot(data,out);document(data,analysis,out);print(out)
