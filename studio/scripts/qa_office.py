#!/usr/bin/env python3
"""Independent numerical / reference / pagination audit of finished Office files."""
import argparse,ast,json,math,re,statistics
from collections import Counter
from pathlib import Path
import openpyxl
from openpyxl.utils.cell import range_boundaries
from docx import Document
import fitz

def evaluate_workbook(path,source):
    formula=openpyxl.load_workbook(path,data_only=False);cached=openpyxl.load_workbook(path,data_only=True)
    findings=[];results=[];active=set();memo={}
    def value(sheet,coord):
        key=(sheet,coord)
        if key in memo:return memo[key]
        if key in active:raise ValueError('Circular reference '+str(key))
        active.add(key);v=formula[sheet][coord].value
        result=expression(v[1:],sheet) if isinstance(v,str) and v.startswith('=') else v
        active.remove(key);memo[key]=result;return result
    def expression(expr,sheet):
        expr=expr.strip()
        if expr.startswith('IFERROR(') and expr.endswith(')'):
            content=expr[8:-1];level=0;split=None
            for i,ch in enumerate(content):
                level+=1 if ch=='(' else -1 if ch==')' else 0
                if ch==',' and level==0:split=i;break
            if split is None:raise ValueError('Invalid IFERROR')
            try:return expression(content[:split],sheet)
            except ZeroDivisionError:return expression(content[split+1:],sheet)
        match=re.fullmatch(r'(SUM|COUNT|AVERAGE|MEDIAN|MIN|MAX)\((?:(?:\x27([^\x27]+)\x27|([A-Za-z_][\w ]*))!)?(\$?[A-Z]+\$?\d+:\$?[A-Z]+\$?\d+)\)',expr)
        if match:
            fn,quoted,plain,rng=match.groups();s=quoted or plain or sheet
            left,top,right,bottom=range_boundaries(rng)
            if (right-left+1)*(bottom-top+1)>10000:raise ValueError('Range exceeds audit budget')
            cells=[value(s,formula[s].cell(r,c).coordinate) for r in range(top,bottom+1) for c in range(left,right+1)]
            nums=[v for v in cells if isinstance(v,(int,float)) and not isinstance(v,bool)]
            return {'SUM':lambda:sum(nums),'COUNT':lambda:len(nums),'AVERAGE':lambda:sum(nums)/len(nums),'MEDIAN':lambda:statistics.median(nums),'MIN':lambda:min(nums),'MAX':lambda:max(nums)}[fn]()
        def ref(m):
            target=m.group(1) or sheet;coord=m.group(2).replace('$','')
            v=value(target,coord)
            if not isinstance(v,(int,float)):raise ValueError('Non-numeric cell '+target+'!'+coord)
            return repr(v)
        replaced=re.sub(r'(?:(\w+)!)?(\$?[A-Z]+\$?\d+)',ref,expr)
        tree=ast.parse(replaced,mode='eval')
        def calc(n):
            if isinstance(n,ast.Expression):return calc(n.body)
            if isinstance(n,ast.Constant) and isinstance(n.value,(int,float)):return n.value
            if isinstance(n,ast.UnaryOp) and isinstance(n.op,(ast.UAdd,ast.USub)):return calc(n.operand)*(1 if isinstance(n.op,ast.UAdd) else -1)
            if isinstance(n,ast.BinOp) and isinstance(n.op,(ast.Add,ast.Sub,ast.Mult,ast.Div)):
                left,right=calc(n.left),calc(n.right)
                if isinstance(n.op,ast.Add):return left+right
                if isinstance(n.op,ast.Sub):return left-right
                if isinstance(n.op,ast.Mult):return left*right
                return left/right
            raise ValueError('Unsupported formula syntax')
        return calc(tree)
    for sheet in formula:
        for row in sheet:
            for cell in row:
                if not (isinstance(cell.value,str) and cell.value.startswith('=')):continue
                try:
                    expected=value(sheet.title,cell.coordinate);observed=cached[sheet.title][cell.coordinate].value
                    ok=isinstance(observed,(int,float)) and math.isclose(expected,observed,rel_tol=1e-10,abs_tol=1e-8)
                    results.append({'cell':sheet.title+'!'+cell.coordinate,'formula':cell.value,'computed':expected,'cached':observed,'status':'PASS' if ok else 'FAIL'})
                    if not ok:findings.append({'code':'FORMULA_CACHE','cell':cell.coordinate,'expected':expected,'actual':observed})
                except (ValueError,KeyError,TypeError,SyntaxError,ZeroDivisionError) as e:findings.append({'code':'FORMULA_NOT_VERIFIED','cell':sheet.title+'!'+cell.coordinate,'error':str(e)})
    for col,key in [('B','control'),('C','assisted')]:
        actual=[formula['Raw'][f'{col}{r}'].value for r in range(2,8)]
        if actual!=source[key]:findings.append({'code':'SOURCE_MISMATCH','column':col})
    expected_medians=[statistics.median(source[k]) for k in ('control','assisted')]
    actual_medians=[cached['Analysis'][x].value for x in ['B5','C5']]
    if actual_medians!=expected_medians:findings.append({'code':'MEDIAN_SOURCE','expected':expected_medians,'actual':actual_medians})
    charts=[]
    for sheet in formula:
        for chart in sheet._charts:
            for ser in chart.series:
                ref=ser.val.numRef.f
                rng=ref.rsplit('!',1)[-1].replace('$','');left,top,right,bottom=range_boundaries(rng)
                if top<1 or left<1 or bottom>sheet.max_row or right>sheet.max_column:findings.append({'code':'CHART_RANGE','reference':ref})
                charts.append({'sheet':sheet.title,'values_reference':ref,'axis_min':chart.y_axis.scaling.min})
                if chart.y_axis.scaling.min!=0:findings.append({'code':'CHART_BASELINE','reference':ref})
    if not charts:findings.append({'code':'MISSING_NATIVE_CHART'})
    return {'status':'FAIL' if findings else 'PASS','scope':'Supported arithmetic and aggregate formulas, cached values, source rows, chart references and zero baseline; not source truth','formulas':results,'charts':charts,'findings':findings}

def document(docx,pdf):
    d=Document(docx);p=fitz.open(pdf);findings=[]
    expected=' '.join([x.text for x in d.paragraphs]+[c.text for t in d.tables for row in t.rows for c in row.cells]);actual=' '.join(page.get_text() for page in p)
    normalize=lambda s:re.sub(r'\W+','',s.lower())
    e,a=Counter(normalize(expected)),Counter(normalize(actual));coverage=sum((e&a).values())/max(1,sum(e.values()))
    if coverage<.97:findings.append({'code':'DOCUMENT_TEXT_LOSS','coverage':coverage})
    if len(p)!=3:findings.append({'code':'PAGINATION','expected':3,'actual':len(p)})
    for i,page in enumerate(p):
        for b in page.get_text('dict')['blocks']:
            for line in b.get('lines',[]):
                for span in line['spans']:
                    x0,y0,x1,y1=span['bbox']
                    if min(x0,y0)<-1 or x1>page.rect.width+1 or y1>page.rect.height+1:findings.append({'code':'TEXT_BOUNDS','page':i+1,'text':span['text'][:80]})
    return {'status':'FAIL' if findings else 'PASS','pages':len(p),'tables':len(d.tables),'paragraphs':len(d.paragraphs),'text_coverage':round(coverage,4),'fonts':sorted({font[3] for page in p for font in page.get_fonts()}),'findings':findings,'scope':'Structure, pagination and text export; actual visual pass is recorded separately'}

if __name__=='__main__':
    p=argparse.ArgumentParser();p.add_argument('mode',choices=['xlsx','docx']);p.add_argument('file');p.add_argument('--source');p.add_argument('--pdf');p.add_argument('--out',required=True);a=p.parse_args()
    result=evaluate_workbook(a.file,json.loads(Path(a.source).read_text())) if a.mode=='xlsx' else document(a.file,a.pdf)
    Path(a.out).write_text(json.dumps(result,ensure_ascii=False,indent=2));print(json.dumps(result,ensure_ascii=False,indent=2));raise SystemExit(result['status']=='FAIL')
