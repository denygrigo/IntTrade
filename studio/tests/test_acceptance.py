#!/usr/bin/env python3
"""Adversarial checks for the production pipeline, not tests of visual taste."""
import importlib.util
import json
from pathlib import Path
import subprocess
import sys
import tempfile
import unittest

ROOT=Path(__file__).resolve().parents[2]
SCRIPTS=ROOT/"studio/scripts"
sys.path.insert(0,str(SCRIPTS))
from qa import evidence,deck

class Acceptance(unittest.TestCase):
    def test_bad_qa_fixture(self):
        result=evidence(json.loads((ROOT/"studio/tests/qa_bad.json").read_text()))
        codes={i["code"] for i in result["findings"]}
        self.assertTrue({"UNSUPPORTED_CLAIM","INCONSISTENT_NUMBER","BROKEN_SOURCE","CALCULATION"}<=codes)
        self.assertEqual(result["status"],"FAIL")

    def test_good_qa_fixture(self):
        self.assertEqual(evidence(json.loads((ROOT/"studio/tests/qa_good.json").read_text()))["status"],"PASS")

    def test_artifact_defects(self):
        from pptx import Presentation
        from pptx.util import Inches,Pt
        with tempfile.TemporaryDirectory() as t:
            r=Presentation();s=r.slides.add_slide(r.slide_layouts[6])
            b=s.shapes.add_textbox(Inches(-1),Inches(1),Inches(3),Inches(1))
            run=b.text_frame.paragraphs[0].add_run();run.text="Intentionally clipped";run.font.size=Pt(6)
            r.slides.add_slide(r.slide_layouts[6])
            f=Path(t)/"broken.pptx";r.save(f)
            result=deck(f);codes={i["code"] for i in result["findings"]}
            self.assertTrue({"OFF_CANVAS","TINY_TYPE","EMPTY_SLIDE"}<=codes)

    def test_recovery_and_hashes(self):
        with tempfile.TemporaryDirectory() as t:
            d=Path(t);out=d/"artifact.txt";source=d/"input.txt";source.write_text("v1")
            base=[sys.executable,str(SCRIPTS/"run_guard.py"),"--state-dir",str(d/"state"),"--stage","build","--input",str(source),"--output",str(out)]
            command=["--",sys.executable,"-c",f"from pathlib import Path;Path({str(out)!r}).write_text('built')"]
            def run(args):return subprocess.run(args,capture_output=True,text=True)
            self.assertEqual(run(base+command).returncode,0)
            self.assertIn("SKIP",run(base+command).stdout)
            out.write_text("tampered");self.assertIn("SUCCESS",run(base+command).stdout)
            fail=["--",sys.executable,"-c","raise SystemExit(1)"]
            self.assertEqual(run(base+fail).returncode,1)
            self.assertEqual(run(base+fail).returncode,1)
            self.assertEqual(run(base+fail).returncode,78)
            self.assertEqual(run(base+["--reset-reason","Changed hypothesis after diagnosis"]+fail).returncode,1)

    def test_timeout_and_missing_output(self):
        with tempfile.TemporaryDirectory() as t:
            base=[sys.executable,str(SCRIPTS/"run_guard.py"),"--state-dir",t,"--stage","timeout","--output",str(Path(t)/"missing"),"--timeout","0.15"]
            r=subprocess.run(base+["--",sys.executable,"-c","import time;time.sleep(3)"],capture_output=True)
            self.assertEqual(r.returncode,124)
            r=subprocess.run(base+["--",sys.executable,"-c","pass"],capture_output=True)
            self.assertEqual(r.returncode,66)
            import fcntl
            with (Path(t)/'busy.lock').open('w') as lock:
                fcntl.flock(lock,fcntl.LOCK_EX|fcntl.LOCK_NB)
                busy=[sys.executable,str(SCRIPTS/'run_guard.py'),'--state-dir',t,'--stage','busy','--output',str(Path(t)/'missing'),'--',sys.executable,'-c','pass']
                self.assertEqual(subprocess.run(busy,capture_output=True).returncode,75)

if __name__=="__main__":unittest.main(verbosity=2)
